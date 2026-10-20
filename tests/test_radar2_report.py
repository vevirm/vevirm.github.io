"""Regression tests for Radar 2's unambiguous Word/audit contract."""
import json
import importlib
from pathlib import Path

from docx import Document

from scripts import retrospective_radar, outcome_verification, outlook_summary


def test_word_separates_candidate_from_observed_outcome(tmp_path, monkeypatch):
    """Both directions are present, but forecasts are not claimed verified."""
    monkeypatch.setattr(retrospective_radar, 'ROOT', tmp_path)
    monkeypatch.setattr(retrospective_radar, 'REPORT', tmp_path/'stuff'/'WHAT_WAS_EXPECTED_WHAT_HAPPENED.docx')
    imported_verification = importlib.import_module('outcome_verification')
    imported_outlook = importlib.import_module('outlook_summary')
    monkeypatch.setattr(imported_verification, 'QUEUE', tmp_path/'outcome-review-queue.json')
    monkeypatch.setattr(imported_verification, 'REGISTRY', tmp_path/'verified-outcomes.json')
    monkeypatch.setattr(imported_outlook, 'ROOT', tmp_path)
    monkeypatch.setattr(imported_outlook, 'OUT', tmp_path/'outlook-summary.json')
    (tmp_path/'research-leads.json').write_text(json.dumps({'leads': []}))
    (tmp_path/'verified-outcomes.json').write_text(json.dumps({'cases': []}))
    (tmp_path/'outcome-observations.json').write_text(json.dumps({'observations': [{
        'topic': 'Emissions', 'expected': 'Policy target by 2030',
        'happened': 'Measured progress reported for 2024.', 'verdict': 'Target not due',
        'sources': [{'date': '2025-01-01', 'finding': 'Official measurement', 'url': 'https://example.org/source'}]
    }]}))
    item = {'id': 'lead-123', 'title': 'Growth forecast', 'source_url': 'https://example.org/original',
            'screening_status': 'candidate_needs_validation',
            'candidate_excerpt': 'Employment will rise by 10% by 2030.', 'candidate_year': 2030,
            'original_publication': '2020-01-01', 'status': 'not_yet_assessed', 'evidence': []}
    retrospective_radar.build_word([item], run={
        'provider': 'evidence', 'attempts': 0, 'failed': 0,
        'new_evidence_assessments': 0, 'note': 'Offline'} )
    text = '\n'.join(p.text for p in Document(retrospective_radar.REPORT).paragraphs)
    assert 'Study comparisons / research attempts this run: 0' in text
    assert 'Source-reviewed expected-vs-actual outcomes: 0' in text
    assert 'WHAT WAS EXPECTED (unverified abstract excerpt): Employment will rise' in text
    assert 'WHAT ACTUALLY HAPPENED: Not established by independent' in text
    assert 'WHAT WAS EXPECTED / TARGETED: Policy target by 2030' in text
    assert 'WHAT ACTUALLY HAPPENED: Measured progress reported for 2024' in text
    assert json.loads((tmp_path/'outcome-review-queue.json').read_text())['candidate_count'] == 1


def test_radar2_scheduled_only_twice_with_secret_and_auto_provider():
    root = Path(__file__).resolve().parents[1]
    workflow = (root/'.github/workflows/update-outcomes.yml').read_text()
    assert "17 4,16 * * *" in workflow
    assert 'workflow_run:' not in workflow, 'Workflow-run plus cron triggers 4 times/day, not 2'
    assert 'OPENAI_API_KEY: ${{ secrets.OPENAI_API_KEY }}' in workflow
    assert '--provider auto' in workflow
    assert 'stuff/WHAT_WAS_EXPECTED_WHAT_HAPPENED.docx' in workflow



def test_word_includes_provisional_happened_when_research_found(tmp_path, monkeypatch):
    import importlib
    from scripts import retrospective_radar
    monkeypatch.setattr(retrospective_radar, 'ROOT', tmp_path)
    monkeypatch.setattr(retrospective_radar, 'REPORT', tmp_path/'stuff'/'WHAT_WAS_EXPECTED_WHAT_HAPPENED.docx')
    verification = importlib.import_module('outcome_verification')
    outlook = importlib.import_module('outlook_summary')
    monkeypatch.setattr(verification, 'QUEUE', tmp_path/'outcome-review-queue.json')
    monkeypatch.setattr(verification, 'REGISTRY', tmp_path/'verified-outcomes.json')
    monkeypatch.setattr(outlook, 'ROOT', tmp_path)
    monkeypatch.setattr(outlook, 'OUT', tmp_path/'outlook-summary.json')
    (tmp_path/'verified-outcomes.json').write_text(json.dumps({'cases': []}))
    (tmp_path/'outcome-observations.json').write_text(json.dumps({'observations': []}))
    item = {'id': 'lead-123', 'title': 'Employment changes in Europe',
            'source_url': 'https://example.org/first', 'original_publication': '2020',
            'screening_status': 'candidate_needs_validation',
            'candidate_excerpt': 'Employment will increase by 10% by 2030.',
            'candidate_year': 2030, 'status': 'early_signals',
            'what_happened': 'Employment increased in the observed dataset as of 2025.',
            'checked_utc': '2026-10-20T12:00:00Z',
            'evidence': [{'url': 'https://example.org/statistics', 'date': '2025',
                          'finding': 'Employment increased in the 2025 release'}]}
    retrospective_radar.build_word([item], run={'provider': 'openai', 'attempts': 1,
                                                 'failed': 0, 'new_evidence_assessments': 1})
    text = '\n'.join(p.text for p in Document(retrospective_radar.REPORT).paragraphs)
    assert 'WHAT ACTUALLY HAPPENED (automated research; NOT source-reviewed): Employment increased' in text
    assert 'Source-reviewed expected-vs-actual outcomes: 0' in text
    assert 'Potential supporting lead: Employment increased' in text


def test_workflow_uploads_fresh_docx_before_research_status_gate():
    root = Path(__file__).resolve().parents[1]
    workflow = (root/'.github/workflows/update-outcomes.yml').read_text()
    assert 'actions/upload-artifact@v4' in workflow
    assert 'name: Radar-2-Word-and-Audit' in workflow
    assert 'if-no-files-found: error' in workflow
    assert 'rm -f stuff/WHAT_WAS_EXPECTED_WHAT_HAPPENED.docx' in workflow
    assert 'python scripts/validate_radar2_report.py' in workflow
    assert workflow.index('Upload Radar 2 Word file') < workflow.index('Enforce live research status')
    assert 'OPENAI_API_KEY: ${{ secrets.OPENAI_API_KEY }}' in workflow
    assert 'allow_offline_report' in workflow
