"""Radar 2 live/offline integrity checks and .docx delivery smoke tests."""
import json
from pathlib import Path
import pytest
from docx import Document
from scripts.check_radar2_live_status import check_live_status
from scripts.validate_radar2_report import validate


def test_no_live_provider_cannot_turn_run_green_by_default():
    audit = {"run": {"provider": "evidence", "attempts": 0, "failed": 0}}
    with pytest.raises(ValueError, match="No live research"):
        check_live_status(audit)
    assert "Offline Word generation" in check_live_status(audit, allow_offline=True)


def test_api_failure_cannot_turn_run_green():
    with pytest.raises(ValueError, match="failed or was skipped"):
        check_live_status({"run": {"provider": "openai", "attempts": 2, "failed": 2}})
    assert "Live research executed" in check_live_status({"run": {"provider": "openai", "attempts": 2, "failed": 1}})


def test_validation_rejects_bad_docx(tmp_path):
    bad = tmp_path / "report.docx"
    bad.write_bytes(b"not a Word file")
    audit = tmp_path / "audit.json"
    audit.write_text(json.dumps({"run": {}, "reviewed_verdicts": 0, "research_leads": 1}))
    with pytest.raises(ValueError, match="missing or too small"):
        validate(bad, audit)


def test_validation_accepts_formatted_radar2_report(tmp_path):
    report = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading('WORLD RADAR 2 | Expected vs. Actually Happened', 0)
    for line in ('Research leads indexed', 'WHAT WAS EXPECTED',
                 'WHAT ACTUALLY HAPPENED', 'Source-reviewed expected-vs-actual outcomes'):
        doc.add_paragraph(line)
    # Add enough content for a realistic, nonempty report.
    doc.add_paragraph('Sample data. ' * 400)
    doc.save(report)
    audit = tmp_path / "audit.json"
    audit.write_text(json.dumps({"run": {}, "reviewed_verdicts": 0, "research_leads": 1}))
    assert validate(report, audit)["research_leads"] == 1
