"""Radar 2: source-grounded, clearly *informal* retrospective judgments.

Uses OpenAlex scholarly metadata by default (no OpenAI dependency). Compares
later studies with the original research and labels all conclusions provisional.
"""
import argparse
import http.client
import random
import socket
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import re
import time
import urllib.error
import urllib.parse
import urllib.request
try:
    from forecast_screening import screen
except ModuleNotFoundError:
    from scripts.forecast_screening import screen
try:
    from openalex_comparison import request_works, related, local_catalogue, assess, query_for
except ModuleNotFoundError:
    from scripts.openalex_comparison import request_works, related, local_catalogue, assess, query_for

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'retrospective-results.json'
REPORT = ROOT / 'stuff' / 'WHAT_WAS_EXPECTED_WHAT_HAPPENED.docx'
MODEL = os.getenv('RADAR_OUTCOME_MODEL', 'gpt-4.1-mini')


def read_json(path, default):
    try:
        return json.loads(path.read_text(encoding='utf-8'))
    except FileNotFoundError:
        return default


def concise(sentence):
    return re.sub(r'\s+', ' ', str(sentence if sentence is not None else '')).strip()


def claims():
    """Every Scanner 1 research lead is eligible for Scanner 2 follow-up.

    A lead is a discovery record, not an authenticated prediction. Do not invent
    a forecast or treat a scenario as a prediction.
    """
    out = {}
    leads = read_json(ROOT / 'research-leads.json', {}).get('leads', [])
    for lead in leads:
        screening = screen(lead)
        url = concise(lead.get('url'))
        title = concise(lead.get('title'))
        if not url or not title:
            continue
        identifier = 'lead-' + hashlib.sha256(url.lower().rstrip('/').encode()).hexdigest()[:16]
        out[identifier] = {'id': identifier, 'expected': concise(lead.get('abstract')) or title,
                           'source_url': url, 'title': title, 'target_year': None,
                           'original_publication': lead.get('date') or 'Unknown',
                           'topic': lead.get('journal') or 'Foresight research',
                           'publisher': lead.get('publisher') or lead.get('journal') or '',
                           'lead_status': lead.get('status') or 'Discovered lead',
                           'is_lead': True, **screening}
    return list(out.values())


def response_text(payload):
    chunks = []
    for output in payload.get('output', []):
        if output.get('type') == 'message':
            for content in output.get('content', []):
                if content.get('type') == 'output_text':
                    chunks.append(content.get('text', ''))
    return '\n'.join(chunks)


# Different evidence paths are deliberately explored on successive runs.
# These are search *directions*, not claims that a given site contains the answer.
EVIDENCE_ROUTES = (
    ('EU data', 'Eurostat, data.europa.eu, ECB Statistical Data Warehouse, EU agencies'),
    ('Member-state data', 'national statistical offices, ministries, central banks, national regulators'),
    ('Global datasets', 'OECD Data, World Bank DataBank, IMF World Economic Outlook, UNdata'),
    ('Peer-reviewed follow-up', 'Crossref, OpenAlex, Google Scholar, scholarly longitudinal studies and replication papers'),
    ('EU policy evaluations', 'European Commission ex-post evaluations, JRC studies, European Court of Auditors'),
    ('Sector and regulator evidence', 'sector regulators, energy agencies, EEA, Eurofound, Cedefop, audited industry series'),
    ('Policy and law', 'EUR-Lex, parliamentary records, official implementation and enforcement reports'),
    ('Historical archives', 'Wayback archival material, historical annual reports, archived time-series releases'),
    ('Energy and climate', 'European Environment Agency, IEA, ENTSO-E, Copernicus and emissions inventories'),
    ('Labour and demography', 'Eurostat labour-force surveys, ILOSTAT, UN population data, national censuses'),
    ('Economy and finance', 'ECB, OECD Economic Outlook, Eurostat national accounts, national central banks'),
    ('Technology and digital', 'DESI/Digital Decade monitoring, OECD ICT statistics, patent and adoption datasets'),
    ('Health and society', 'WHO Europe, ECDC, OECD Health at a Glance, national health authorities'),
    ('Agriculture and biodiversity', 'FAOSTAT, EEA biodiversity reports, Copernicus land monitoring, Eurostat agriculture'),
    ('Transport and mobility', 'Eurostat transport statistics, ACEA, national registration agencies, EU transport observatories'),
    ('Contradictory evidence', 'independent counter-evaluations, methodological critiques, rival time series and uncertainty ranges'),
)

# Per-claim research alternates routes across runs. These are starting points,
# not a closed whitelist: the agent is told to search beyond them if necessary.


def route_for(entry):
    """Cycle evidence sources for each claim without repeatedly trying one route."""
    attempts = max(0, int(entry.get('research_attempts', 0)))
    # A stable claim-specific offset spreads source categories across a single run.
    offset = int(hashlib.sha256(str(entry.get('id', '')).encode()).hexdigest()[:8], 16)
    return (offset + attempts) % len(EVIDENCE_ROUTES)


def retrieve_assessment(claim, key, route=0):
    now = dt.date.today().isoformat()
    route_name, route_sources = EVIDENCE_ROUTES[route % len(EVIDENCE_ROUTES)]
    prompt = f'''Research route: {route_name}. Start with {route_sources}, then broaden as needed. Today: {now}.
Investigate this European foresight research LEAD and identify any subsequent real-world developments relevant to its themes. A title or abstract may describe a scenario, not a prediction. Do not invent a forecast.  Use WEB SEARCH for independently dated evidence AFTER the original publication. Search for both positive signals and counterevidence. Distinguish actual deployment or measured change from investment announcements, plans, ambitions, scenarios, and publicity. A future target date DOES NOT prevent assessment of current signals; it only prevents calling the final outcome realized prematurely.
Research lead title/abstract (NOT necessarily a prediction): {claim['expected']}
Original study: {claim['source_url']}
Original publication: {claim.get('original_publication')}
Target year: {claim.get('target_year')}

Return ONLY one JSON object with:
status: exactly one of realized, materializing, early_signals, no_clear_evidence, contradicted, too_early, insufficient_evidence.
what_expected: faithful concise description.
what_happened: concise observed developments, or clearly say none established.
evidence: list of objects with url, date, finding (dated independent sources; preferably two if feasible).
counterevidence: list of objects with url, date, finding when found.
reasoning: how evidence supports status, limitations and differences in baselines/time/geography.
watch_next: specific observable developments or measurable milestones to check.

Rules: 'realized' requires clear completed outcome with direct evidence; 'materializing' requires significant observed implementation; 'early_signals' requires credible observable indications, not mere publicity; 'contradicted' requires direct material counterevidence; 'no_clear_evidence' is only for a genuine search without convincing signals; 'too_early' means a judgment is premature AND there are no present signals sufficient to classify; 'insufficient_evidence' means research failed to establish reliable support. Future scenario narratives should not be treated as promised predictions. Never fabricate citations. Do not cite the original publication as the ONLY later evidence.'''
    request = {'model': MODEL, 'tools': [{'type': 'web_search'}],
               'input': prompt, 'max_output_tokens': 1800}
    req = urllib.request.Request('https://api.openai.com/v1/responses',
        data=json.dumps(request).encode('utf-8'),
        headers={'Authorization': 'Bearer ' + key, 'Content-Type': 'application/json'}, method='POST')
    with urllib.request.urlopen(req, timeout=100) as response:
        data = json.load(response)
    raw = response_text(data)
    match = re.search(r'\{[\s\S]*\}', raw)
    if not match:
        raise ValueError('No structured assessment in response')
    return json.loads(match.group())


def free_follow_up(claim, timeout=12):
    """Look for subsequent scholarship in OpenAlex (free; not outcome verification).

    A discovered paper is a research follow-up, NOT evidence a scenario happened.
    """
    original = claim.get('original_publication', '')[:4]
    year = int(original) if original.isdigit() else 0
    if not year:
        return []
    title = concise(claim['title'])[:150]
    if not title:
        return []
    # A simpler OpenAlex query avoids invalid select/filter combinations on
    # heterogeneous records. Filter the dates locally, including future years.
    params = {'search': title, 'per_page': 5}
    email = os.getenv('MAILTO', '').strip() or os.getenv('OPENALEX_EMAIL', '').strip()
    if email and '@' in email:
        params['mailto'] = email
    api_key = os.getenv('OPENALEX_API_KEY', '').strip()
    if api_key:
        params['api_key'] = api_key
    def fetch_works(query_params):
        """Retry transport failures, but never retry authorization/quota failures blindly."""
        url = 'https://api.openalex.org/works?' + urllib.parse.urlencode(query_params)
        req = urllib.request.Request(url, headers={
            'User-Agent': f'WorldRadar/3.1 (mailto:{email})' if '@' in email else 'WorldRadar/3.1',
            'Accept': 'application/json', 'Accept-Encoding': 'identity'})
        for attempt in range(4):
            try:
                with urllib.request.urlopen(req, timeout=timeout) as response:
                    # Read the entire response inside the protected retry block.
                    payload = response.read()
                data = json.loads(payload)
                if not isinstance(data, dict) or not isinstance(data.get('results'), list):
                    raise ValueError('OpenAlex returned an unexpected results payload')
                return data['results']
            except urllib.error.HTTPError as exc:
                # A 400 may be fixed by the simplified search below. A 429 may
                # indicate daily quota exhaustion: do not spend requests retrying.
                if exc.code not in (408, 500, 502, 503, 504):
                    raise
            except (http.client.IncompleteRead, http.client.RemoteDisconnected,
                    ConnectionResetError, BrokenPipeError, TimeoutError,
                    socket.timeout, urllib.error.URLError) as exc:
                print(f'OpenAlex transient transport error ({type(exc).__name__}); retry {attempt + 1}/4', flush=True)
            if attempt == 3:
                break
            time.sleep(min(8, 1.5 * (2 ** attempt)) + random.uniform(0, 0.3))
        raise ConnectionError('OpenAlex response incomplete or unavailable after 4 attempts')

    try:
        items = fetch_works(params)
    except urllib.error.HTTPError as exc:
        if exc.code != 400:
            raise
        fallback = {'search': ' '.join(title.split()[:8]), 'per_page': 5}
        if email and '@' in email:
            fallback['mailto'] = email
        if api_key:
            fallback['api_key'] = api_key
        items = fetch_works(fallback)
    stopwords = {'study','research','scenarios','europe','european','future','towards','policy','analysis','perspective','development','modelling'}
    def tokens(t):
        return {w for w in re.findall(r'[a-z]{4,}', t.lower()) if w not in stopwords}
    original_tokens = tokens(title)
    found = []
    for work in items:
        if (work.get('publication_year') or 0) <= year:
            continue
        wtitle = concise(work.get('title'))
        if not wtitle or wtitle.lower() == title.lower():
            continue
        terms = tokens(wtitle)
        overlap = len(terms & original_tokens)
        # Conservative relatedness: do not treat generic papers as follow-ups.
        if overlap < 3 or overlap / max(1, min(len(terms), len(original_tokens))) < .35:
            continue
        link = work.get('doi') or work.get('id')
        if not link or not link.startswith('https://') or link.rstrip('/') == claim['source_url'].rstrip('/'):
            continue
        found.append({'title': wtitle, 'url': link, 'year': work.get('publication_year'),
                      'note': 'Potentially related later publication; outcome not independently verified'})
        if len(found) >= 3:
            break
    return found


def single_sentence(value):
    value = concise(value)
    if not value:
        return ''
    # Avoid a multi-sentence narrative where the Word report requires one sentence.
    return re.split(r'(?<=[.!?])\s+(?=[A-Z])', value, maxsplit=1)[0].rstrip('.!?') + '.'


STATUS_LABELS = {
    'realized': 'Realized', 'materializing': 'Materializing',
    'early_signals': 'Early signals', 'no_clear_evidence': 'No clear evidence',
    'contradicted': 'Contradicted', 'too_early': 'Too early to tell',
    'insufficient_evidence': 'Insufficient evidence',
    'not_yet_assessed': 'Not yet investigated',
}


def clean_evidence(claim, values):
    sources = []
    for e in values if isinstance(values, list) else []:
        if not isinstance(e, dict):
            continue
        url = str(e.get('url', '')).strip()
        if (not re.match(r'^https?://[^\s/]+', url)
                or url.rstrip('/') == claim['source_url'].rstrip('/')
                or not e.get('finding') or not e.get('date')):
            continue
        sources.append({'url': url, 'date': str(e['date'])[:40], 'finding': concise(e['finding'])[:500]})
    return sources[:8]


def normalize(claim, result):
    status = result.get('status', 'insufficient_evidence')
    if status not in STATUS_LABELS or status == 'not_yet_assessed':
        status = 'insufficient_evidence'
    evidence = clean_evidence(claim, result.get('evidence'))
    counter = clean_evidence(claim, result.get('counterevidence'))
    # No definite conclusion without independently dated supporting evidence.
    if status in ('realized', 'materializing', 'early_signals', 'contradicted') and not evidence and not counter:
        status = 'insufficient_evidence'
    target = claim.get('target_year')
    if target and int(target) > dt.date.today().year and status == 'realized':
        # Only allow realization if the research actually demonstrates early completion.
        if not evidence:
            status = 'insufficient_evidence'
    return {**claim, 'what_expected': concise(result.get('what_expected') or claim['expected']),
            'what_happened': concise(result.get('what_happened') or 'No independently established observations yet.'),
            'status': status, 'evidence': evidence, 'counterevidence': counter,
            'reasoning': concise(result.get('reasoning', '')),
            'watch_next': concise(result.get('watch_next', '')),
            'checked_utc': dt.datetime.now(dt.timezone.utc).isoformat()}


def build_word(entries, run=None):
    """Generate the Radar 2 Word report from auditable saved data.

    Do not confuse potential forecasts, model-assisted leads or policy targets
    with source-reviewed realized outcomes. The report is useful offline too.
    """
    from collections import Counter
    from docx import Document
    from docx.shared import Inches, Pt
    from outlook_summary import create
    from outcome_verification import assemble, write_queue

    create(entries)
    verified, rejected = assemble(entries)
    write_queue(entries)
    candidates = sorted((e for e in entries if e.get('screening_status') == 'candidate_needs_validation'),
                        key=lambda e: (e.get('candidate_year') or 9999, e.get('title', '')))
    screening = Counter(e.get('screening_status', 'unscreened') for e in entries)
    run = run or {}
    observations = read_json(ROOT / 'outcome-observations.json', {}).get('observations', [])
    now = dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%d %H:%M UTC')

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    for name in ('Heading 1', 'Heading 2'):
        doc.styles[name].paragraph_format.space_before = Pt(11)
        doc.styles[name].paragraph_format.space_after = Pt(5)

    def field(label, value):
        paragraph = doc.add_paragraph()
        paragraph.add_run(label + ': ').bold = True
        paragraph.add_run(concise(value) if value is not None and str(value).strip() else 'Not established.')
        return paragraph

    doc.add_heading('WORLD RADAR 2 | Expected vs. Actually Happened', 0)
    doc.add_paragraph('Evidence audit generated ' + now)
    doc.add_paragraph(f'Important: {len(entries)} discovery leads are not {len(entries)} verified predictions. A report regenerated today does not imply new evidence was discovered today. Forecast candidates, observed policy developments, and reviewed outcomes are separate categories.')
    doc.add_heading('Run and evidence summary', 1)
    field('Research leads indexed', len(entries))
    field('Potential forecast excerpts (not validated)', len(candidates))
    field('Candidate horizons still ahead', sum(bool(e.get('candidate_year')) and int(e['candidate_year']) > dt.date.today().year for e in candidates))
    field('Missing full-text prediction checks', screening['full_text_required'])
    field('No dated, testable forecast in metadata', screening['no_testable_prediction_in_metadata'])
    field('Source-reviewed expected-vs-actual outcomes', len(verified))
    field('Studies with tentative OpenAlex/corpus comparisons', sum(bool(e.get('literature_assessed_utc')) for e in entries))
    field('Comparisons with later abstract-reported study results', sum(e.get('literature_signal') == 'later_studies_report_results' for e in entries))
    field('Rejected submitted outcome cases', len(rejected))
    if run:
        field('Run type', run.get('provider', 'unknown'))
        field('Study comparisons / research attempts this run', run.get('attempts', 0))
        field('Live external research requests', run.get('attempts', 0) if run.get('provider') in ('openalex', 'openai') else 0)
        field('Failed research requests', run.get('failed', 0))
        field('New dated-evidence assessments this run', run.get('new_evidence_assessments', 0))
        field('New research comparisons this run', run.get('literature_comparisons', 0))
        field('Later studies found this run', run.get('literature_studies_found', 0))
        field('Run note', run.get('note', ''))
    if not verified:
        doc.add_paragraph('RESULT: No original-text-reviewed, independently measured forecast comparisons have been verified. Any blank outcome below is a genuine evidence gap, not a negative forecast verdict.')

    if verified:
        doc.add_heading('1 | Source-reviewed forecast outcomes', 1)
        for case in verified:
            original, observed = case['original'], case['observation']
            title = next((e.get('title') for e in entries if e['id'] == case['lead_id']), case['lead_id'])
            doc.add_heading(concise(title)[:140], 2)
            field('WHAT WAS EXPECTED (original verbatim)', original['verbatim_passage'])
            field('Original publication', original['url'] + ' | Published: ' + original['published_on'] + ' | Target: ' + original['target_date'])
            field('WHAT ACTUALLY HAPPENED', observed['what_happened'])
            field('Verdict', case['verdict'].replace('_', ' ').capitalize())
            field('Metric and geography comparison', observed['same_metric_and_geography_explained'])
            field('Measured on', observed['measured_on'])
            field('Limitations', case['limitations'])
            field('Human reviewer', case['source_reviewed_by'] + ' | ' + case['reviewed_on'])
            for source in observed['sources']:
                field('Independent evidence (' + source['date'] + ')', source['finding'] + ' | ' + source['url'])

    doc.add_heading('2 | Documented developments (not forecast verdicts)', 1)
    doc.add_paragraph('These manually provided, source-linked observations show measurable policy progress. They must not be counted as validation of a forecast unless the original publication made that exact testable prediction.')
    if not observations:
        doc.add_paragraph('No separately recorded, editor-reviewed policy observations.')
    for item in observations:
        doc.add_heading(concise(item.get('topic'))[:140], 2)
        field('WHAT WAS EXPECTED / TARGETED', item.get('expected'))
        field('WHAT ACTUALLY HAPPENED', item.get('happened'))
        field('Status and caveat', item.get('verdict'))
        query = concise(item.get('original_lead_title_contains')).lower()
        matching = next((e for e in entries if query and query in e.get('title', '').lower()), None)
        if matching:
            field('Related research lead (not necessarily a prediction)', matching.get('title') + ' | ' + matching.get('source_url'))
        for source in item.get('sources', []):
            field('Dated source (' + concise(source.get('date')) + ')', concise(source.get('finding')) + ' | ' + concise(source.get('url')))

    doc.add_heading('3 | Provisional comparisons with later scholarship', 1)
    doc.add_paragraph('This section provides a first-pass assessment using later studies indexed by OpenAlex or already captured by Scanner 1. It is useful for identifying emerging agreement, empirical follow-up, methodological debate and gaps. Abstract statements are the published authors\' reported findings, NOT independent observation confirmed by Radar 2. The word "EU" here means EU-related subjects, not necessarily foresight issued by an EU institution. No performance score is inferred.')
    studies_checked = [e for e in entries if e.get('literature_assessed_utc')]
    prioritized = sorted(studies_checked, key=lambda e: (
        bool(e.get('follow_up_publications')),
        e.get('literature_signal') == 'later_studies_report_results',
        e.get('screening_status') == 'candidate_needs_validation',
        len(e.get('follow_up_publications') or [])), reverse=True)
    field('Records compared so far', len(studies_checked))
    topic_counts = Counter(e.get('topic_cluster', 'Unclassified') for e in studies_checked)
    doc.add_paragraph('Topic coverage of literature comparisons (NOT EU prediction accuracy): ' +
                      '; '.join(f'{topic}: {count}' for topic, count in topic_counts.most_common(8)))
    field('With result-like passages in later abstracts', sum(e.get('literature_signal') == 'later_studies_report_results' for e in studies_checked))
    field('With only related models or topic papers', sum(e.get('literature_signal') == 'later_research_without_comparable_outcome' for e in studies_checked))
    field('With no sufficiently close study found', sum(e.get('literature_signal') == 'no_comparable_study_found' for e in studies_checked))
    if len(prioritized) > 40:
        doc.add_paragraph('Showing the top 40 comparisons, prioritizing records with later-source matches. All assessments, including failed matches, remain in retrospective-results.json.')
    if not prioritized:
        doc.add_paragraph('No comparison run recorded yet. Enable the OpenAlex workflow or run --provider catalogue for existing indexed studies.')
    for number, entry in enumerate(prioritized[:40], 1):
        doc.add_heading(f'{number}. {concise(entry.get("title"))[:120]}', 2)
        field('WHAT WAS SAID / EXPECTED', entry.get('what_expected_literature') or entry.get('candidate_excerpt') or entry.get('expected'))
        field('LATER STUDIES / WHAT THEY REPORT', entry.get('what_happened_literature'))
        field('TENTATIVE ASSESSMENT', entry.get('tentative_assessment'))
        if entry.get('tentative_direction') not in (None, 'not_a_validated_forecast', 'insufficient_comparability'):
            field('Very tentative directional clue', str(entry.get('tentative_direction')).replace('_', ' '))
        field('Confidence / limitations', entry.get('comparison_limits'))
        field('Topic', entry.get('topic_cluster'))
        field('EU institutional attribution', entry.get('eu_attribution'))
        field('Original research', entry.get('source_url'))
        field('Search method and date', str(entry.get('literature_search_mode')) + ' | ' + str(entry.get('literature_assessed_utc'))[:16])
        for source in (entry.get('follow_up_publications') or [])[:4]:
            field('Later study (' + str(source.get('year')) + '; ' + source.get('study_type', 'study') + ')',
                  source.get('title', '') + ' | ' + source.get('url', ''))
            if source.get('finding_from_abstract'):
                field('Its abstract states', source['finding_from_abstract'])

    doc.add_heading('4 | Forecast passages needing original-source validation', 1)
    doc.add_paragraph('The following are candidates from abstracts only, NOT authenticated predictions. They have no source-reviewed actual-outcome verdict. Source-linked machine research, where available, is recorded as a lead for review only.')
    if not candidates:
        doc.add_paragraph('No forecast candidates extracted from the current metadata.')
    for number, entry in enumerate(candidates, 1):
        doc.add_heading(f'{number}. {concise(entry.get("title"))[:135]}', 2)
        field('WHAT WAS EXPECTED (unverified abstract excerpt)', entry.get('candidate_excerpt'))
        field('Original research URL', entry.get('source_url'))
        field('Possible target year (unconfirmed)', entry.get('candidate_year') or 'Not reliably extracted')
        # An AI-researched evidence lead belongs in the report, but must not be
        # presented as a human-reviewed verification of the original forecast.
        # Previously, even a successful evidence search always printed the same
        # blank 'WHAT ACTUALLY HAPPENED' placeholder for candidate forecasts.
        has_research_evidence = bool(entry.get('evidence') or entry.get('counterevidence'))
        provisional_actual = concise(entry.get('what_happened'))
        if has_research_evidence and entry.get('checked_utc') and provisional_actual:
            field('WHAT ACTUALLY HAPPENED (automated research; NOT source-reviewed)', provisional_actual)
            field('Reliability', 'Provisional evidence lead. Check the linked documents, dates, original prediction and like-for-like metrics before interpreting this as a forecast outcome.')
        else:
            field('WHAT ACTUALLY HAPPENED', 'Not established by independent, source-reviewed comparable evidence.')
            if entry.get('literature_assessed_utc') and entry.get('follow_up_publications'):
                field('WHAT LATER SCHOLARSHIP REPORTS (not verified actual outcome)', entry.get('what_happened_literature'))
                field('Tentative interpretation', entry.get('tentative_assessment'))
        field('Verification status', 'Needs original passage, published date, metric/geography/target and later official observations')
        if entry.get('checked_utc'):
            field('Last automated evidence-search attempt', entry['checked_utc'])
        elif entry.get('last_attempt_utc'):
            field('Last failed automated attempt', entry['last_attempt_utc'])
        else:
            field('Last automated evidence search', 'No successful research logged')
        if entry.get('last_error'):
            field('Last research error', entry['last_error'])
        if entry.get('evidence') or entry.get('counterevidence'):
            field('Unreviewed automated research status', entry.get('status', 'unknown') + ' — not a verified outcome')
            for source in (entry.get('evidence') or [])[:3]:
                field('Potential supporting lead', source.get('finding', '') + ' | ' + source.get('date', '') + ' | ' + source.get('url', ''))
            for source in (entry.get('counterevidence') or [])[:3]:
                field('Potential counterevidence lead', source.get('finding', '') + ' | ' + source.get('date', '') + ' | ' + source.get('url', ''))

    doc.add_heading('5 | Coverage and method', 1)
    doc.add_paragraph(f'{len(entries) - len(candidates)} additional catalogued leads are not listed as forecasts because their metadata did not contain an extractable time-bounded prediction. All {len(entries)} source records are in research-leads.json and retrospective-results.json; the {len(candidates)} candidates and next actions are in outcome-review-queue.json. Only validated comparisons entered in verified-outcomes.json count toward reviewed verdicts.')
    doc.add_paragraph('Automated web research can find useful links but cannot certify the content of a source or the comparability of statistics. A human reviewer must validate original wording, dates, metrics, independent outcome documents and limitations before an outcome is labeled verified.')
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--minutes', type=float, default=18)
    ap.add_argument('--max-claims', type=int, default=515)
    ap.add_argument('--offline', action='store_true')
    ap.add_argument('--provider', choices=['auto', 'openalex', 'catalogue', 'free', 'openai', 'evidence'], default='auto')
    args = ap.parse_args()
    stored = read_json(OUTPUT, {'entries': {}})
    entries = stored.setdefault('entries', {})
    found = claims()
    # Legacy numeric scores are not evidence of realization; migrate safely.
    for entry in entries.values():
        entry.pop('score', None)
        if entry.get('status') == 'assessed':
            entry['status'] = 'not_yet_assessed'
            entry['what_happened'] = 'Legacy numerical judgment removed; awaiting evidence-based review.'
            entry['reasoning'] = ''
        entry.setdefault('counterevidence', [])
    key = os.getenv('OPENAI_API_KEY', '').strip()
    if args.provider == 'auto':
        args.provider = 'catalogue' if args.offline else 'openalex'
    if args.provider == 'free':
        args.provider = 'openalex'  # backwards-compatible flag, now meaningful comparisons
    if args.offline and args.provider == 'openalex':
        args.provider = 'catalogue'
    if args.provider == 'openai' and not args.offline and not key:
        raise SystemExit('OpenAI provider explicitly selected but OPENAI_API_KEY is missing.')
    print(f'Radar 2 start: {len(found)} research leads; provider={args.provider}; OpenAlex API key configured={bool(os.getenv("OPENALEX_API_KEY", "").strip())}', flush=True)
    deadline = time.monotonic() + max(0, args.minutes) * 60
    worked = 0
    completed = 0
    failed = 0
    throttled = False
    successful_evidence = 0
    literature_comparisons = 0
    literature_studies_found = 0
    saved_leads = read_json(ROOT / "research-leads.json", {}).get("leads", [])
    for claim in found:
        if claim['id'] not in entries:
            entries[claim['id']] = {**claim, 'what_expected': single_sentence(claim['expected']),
                                   'what_happened': 'No independently verified outcome yet.',
                                   'status': 'not_yet_assessed', 'evidence': [], 'counterevidence': []}
    # Keep the JSON limited to current candidates: legacy/orphaned entries previously
    # made the stored report appear more complete than its live extraction pipeline.
    active_ids = {claim['id'] for claim in found}
    entries = {key: value for key, value in entries.items() if key in active_ids}
    stored['entries'] = entries
    for claim in found:
        entries[claim['id']].update({k: claim[k] for k in ('screening_status','screening_reason','candidate_excerpt','candidate_year')})
    # Prioritize real forecast candidates before general EU-related scholarship.
    # Rotate on the number of *literature checks*, not failed OpenAI requests.
    def priority(c):
        e = entries[c['id']]
        return (int(e.get('literature_live_attempts') or 0) if args.provider == 'openalex' else int(e.get('literature_attempts') or 0),
                0 if c.get('screening_status') == 'candidate_needs_validation' else 1,
                e.get('literature_assessed_utc', ''), c['id'])
    ordered = [] if args.provider == 'evidence' else sorted(found, key=priority)
    groups = {}
    for claim in ordered:
        host = urllib.parse.urlparse(claim['source_url']).hostname or ''
        groups.setdefault(host, []).append(claim)
    queue = []
    while any(groups.values()):
        for host in sorted(groups):
            if groups[host]:
                queue.append(groups[host].pop(0))
    for claim in queue:
        if worked >= args.max_claims or time.monotonic() >= deadline - 8:
            break
        previous = entries[claim['id']]
        try:
            route = route_for(previous)
            print('Research:', args.provider, claim['id'], flush=True)
            previous['research_attempts'] = int(previous.get('research_attempts') or 0) + 1
            if args.provider in ('openalex', 'catalogue'):
                previous['literature_attempts'] = int(previous.get('literature_attempts') or 0) + 1
                if args.provider == 'openalex':
                    previous['literature_live_attempts'] = int(previous.get('literature_live_attempts') or 0) + 1
                if args.provider == 'catalogue':
                    studies = local_catalogue(claim, saved_leads)
                    research_mode = 'existing_scanner_catalogue (saved OpenAlex/Crossref records)'
                else:
                    query = query_for(claim, previous['literature_live_attempts'] - 1)
                    if len(query) < 4:
                        raise ValueError('Not enough substantive terms for a meaningful OpenAlex query')
                    works = request_works(query, int(claim.get('original_publication','')[:4]) + 1)
                    studies = related(claim, works)
                    # These catalogue items supplement the live query, with a
                    # separate provenance label. They do not make a failed API green.
                    if len(studies) < 2:
                        seen = {x['url'] for x in studies}
                        studies.extend(x for x in local_catalogue(claim, saved_leads) if x['url'] not in seen)
                        studies = studies[:4]
                    research_mode = 'openalex_live + saved catalog (when needed)'
                previous.update(assess(claim, studies, research_mode))
                previous.pop('last_error', None)
                previous['bibliography_checked_utc'] = previous['literature_assessed_utc']
                worked += 1
                literature_comparisons += 1
                literature_studies_found += len(studies)
                print('Compared:', claim['id'], 'later studies:', len(studies),
                      'abstract-result snippets:', sum(bool(x.get('finding_from_abstract')) for x in studies), flush=True)
                if args.provider == 'openalex':
                    time.sleep(.3)
            elif args.provider == 'openai':
                result = retrieve_assessment(claim, key, route=route)
                assessed = normalize(claim, result)
                if assessed.get('evidence') or assessed.get('counterevidence'):
                    successful_evidence += 1
                assessed['research_attempts'] = previous['research_attempts']
                assessed['last_route'] = EVIDENCE_ROUTES[route][0]
                if previous.get('literature_assessed_utc'):
                    for k in ('literature_assessed_utc', 'literature_search_mode', 'literature_signal',
                              'what_expected_literature', 'what_happened_literature', 'tentative_assessment',
                              'comparison_limits', 'eu_attribution', 'follow_up_publications', 'literature_attempts', 'literature_live_attempts', 'tentative_direction', 'topic_cluster'):
                        if k in previous:
                            assessed[k] = previous[k]
                if previous.get('status') in ('realized', 'materializing', 'early_signals', 'contradicted') and assessed['status'] == 'insufficient_evidence':
                    previous['checked_utc'] = assessed['checked_utc']
                    previous['last_route'] = assessed['last_route']
                else:
                    entries[claim['id']] = assessed
                worked += 1
                if entries[claim['id']]['status'] in ('realized', 'materializing', 'early_signals', 'contradicted'):
                    completed += 1
                print('Assessed:', claim['id'], entries[claim['id']]['status'], flush=True)
        except (urllib.error.URLError, ValueError, KeyError, json.JSONDecodeError, TimeoutError, ConnectionError, http.client.IncompleteRead, http.client.RemoteDisconnected, ConnectionResetError) as exc:
            print('Assessment failed; prior result kept:', claim['id'], str(exc)[:250], flush=True)
            entries[claim['id']]['last_attempt_utc'] = dt.datetime.now(dt.timezone.utc).isoformat()
            entries[claim['id']]['last_route'] = ('OpenAlex similar-study comparison' if args.provider in ('openalex', 'catalogue') else EVIDENCE_ROUTES[route][0])
            failed += 1
            entries[claim['id']]['last_error'] = str(exc)[:250]
            worked += 1
            OUTPUT.write_text(json.dumps(stored, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
            if isinstance(exc, urllib.error.HTTPError) and exc.code in (401, 403, 429):
                throttled = True
                print('OpenAlex rejected credentials / quota: stopping so future runs can retry.', flush=True)
                break
        if worked and worked % 20 == 0:
            OUTPUT.write_text(json.dumps(stored, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
            print(f'Progress: {worked}/{min(len(found), args.max_claims)} attempts, {failed} failed', flush=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(stored, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    run = {'provider': args.provider, 'attempts': worked, 'failed': failed,
           'new_evidence_assessments': successful_evidence,
           'literature_comparisons': literature_comparisons,
           'literature_studies_found': literature_studies_found,
           'note': ('OpenAlex similar-study comparisons; academic abstracts are provisional, not verified outcomes.'
                    if args.provider == 'openalex' else
                    'Compared genuine existing scholarly records offline; no new external request.'
                    if args.provider == 'catalogue' else
                    'No external research performed.' if args.provider == 'evidence' else
                    'AI-assisted research; citations require original-source checking.')}
    build_word([entries[c['id']] for c in found], run=run)
    from outcome_verification import assemble
    valid_cases, invalid_cases = assemble([entries[c['id']] for c in found])
    audit = {'research_leads': len(found), 'candidate_forecasts': sum(c.get('screening_status') == 'candidate_needs_validation' for c in found), 'reviewed_verdicts': len(valid_cases), 'invalid_cases': invalid_cases, 'outcome_verification_complete': bool(valid_cases), 'run': run, 'literature_compared_total': sum(bool(e.get('literature_assessed_utc')) for e in entries.values()),
             'generated_utc': dt.datetime.now(dt.timezone.utc).isoformat()}
    (ROOT / 'scanner2-audit.json').write_text(json.dumps(audit, indent=2)+'\n', encoding='utf-8')
    print('OUTCOME AUDIT:', json.dumps(audit), flush=True)
    print(f'Radar 2: {len(found)} source records, {worked} bibliography attempts, {failed} failures, {sum(bool(x.get("evidence") or x.get("counterevidence")) for x in entries.values())} entries with dated evidence; Word: {REPORT}', flush=True)
    summary_file = os.getenv('GITHUB_STEP_SUMMARY')
    if summary_file:
        with open(summary_file, 'a', encoding='utf-8') as summary:
            summary.write(f'## Radar 2 follow-up\n\n- Leads: {len(found)}\n- Attempts this run: {worked}\n- Failed this run: {failed}\n- Successful checks saved: {sum(bool(x.get("checked_utc")) for x in entries.values())}\n- OpenAlex throttled: {throttled}\n- Source-reviewed forecast comparisons: {len(valid_cases)}\n- Outcome verification complete: {bool(valid_cases)}\n- Word report generated: {REPORT}\n')
    if not args.offline and args.provider not in ('catalogue', 'evidence') and (worked == 0 or worked == failed) and queue:
        print('::warning::No successful research checks; an honest Word audit was still generated and will be committed.', flush=True)


if __name__ == '__main__':
    main()
